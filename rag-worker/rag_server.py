"""CSIS RAG worker(FastAPI)— 文件解析、chunking、embedding、per-site 向量檢索。

移植自 SuperAI python/rag/rag_server.py,主要改造:
- 多租戶:每個工地一個 ChromaDB collection(site_{siteId}),BM25 索引也 per-site,
  A 工地永遠查不到 B 工地的內容
- 原始檔不落地(CSIS 原檔存 GridFS,由 Bun 端管理),worker 只管索引
- /query 回傳 rerank_score(citation 相關度顯示用)
- /upload 回傳 page_count(文件列表頁數欄位用)

只綁 127.0.0.1(僅供本機 Bun 端呼叫),port 預設 8010。
"""

from __future__ import annotations

import io
import os
import re
import uuid

import chromadb
import uvicorn
from fastapi import FastAPI, Form, HTTPException, UploadFile, File
from markitdown import MarkItDown
from pydantic import BaseModel

# Paginated parsers — fallback to markitdown if these fail at import or runtime.
try:
    import pdfplumber  # type: ignore
except Exception:
    pdfplumber = None  # PDF page-aware parsing will degrade to markitdown

try:
    from pptx import Presentation  # type: ignore
except Exception:
    Presentation = None  # PPTX slide-aware parsing will degrade to markitdown

# Hybrid 檢索:BM25(關鍵字/稀疏)+ jieba 中文分詞。缺套件時降級為純向量檢索。
try:
    import jieba  # type: ignore
    from rank_bm25 import BM25Okapi  # type: ignore
    _BM25_AVAILABLE = True
    jieba.setLogLevel(60)  # 關掉 jieba 啟動時的 logging 噪音
except Exception:
    _BM25_AVAILABLE = False

from embeddings import embed_texts, embed_query, get_embedding_model, get_current_model_name
from log import get_logger

logger = get_logger("RAG")

app = FastAPI(title="CSIS RAG Worker")

# Markitdown converter(DOCX/XLSX 主力,PDF/PPTX 的 fallback)
markitdown = MarkItDown()

MARKITDOWN_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}
PLAINTEXT_EXTENSIONS = {".txt", ".md"}

# ChromaDB client(persistent,單程序使用 — 只有本 worker 碰這個目錄)
CHROMA_DIR = os.environ.get("CHROMA_DIR", os.path.join(os.path.dirname(__file__), "data", "chromadb"))
chroma_client: chromadb.PersistentClient | None = None

# per-site collection cache(get_or_create 有成本,查詢熱路徑不重複走)
_collections: dict[str, object] = {}

# siteId 是 MongoDB ObjectId 的 24 字小寫 hex — 同時是 injection 防護的白名單
_SITE_ID_RE = re.compile(r"[0-9a-f]{24}")

# BM25 稀疏索引(in-memory,per-site,與該 site collection 的 documents 同步)。
# 用「原文 chunk」建(不含《檔名》前綴)→ dense 找對文件、BM25 找對精確 chunk。
# 必須 per-site:全域索引會讓 BM25 這一路把別的工地的 chunk 融進 RRF 結果。
_bm25_by_site: dict[str, tuple[object, list[str]]] = {}

RRF_K = 60              # RRF 慣用常數
HYBRID_CANDIDATES = 50  # 每路召回的候選數(融合前)

# Cross-encoder reranker:召回 → RRF 融合 → 重排。
# 池 8 + 截斷 384 tokens 是無 GPU 機器的平衡值(16 筆全長實測 23-31s 會爆 timeout)。
ENABLE_RERANKER = os.environ.get("ENABLE_RERANKER", "1") == "1"
RERANKER_MODEL = os.environ.get("RERANKER_MODEL", "BAAI/bge-reranker-v2-m3")
RERANK_POOL = int(os.environ.get("RERANK_POOL", "8"))
RERANK_MAX_LENGTH = int(os.environ.get("RERANK_MAX_LENGTH", "384"))
_reranker = None  # CrossEncoder 單例(startup 載入)

# Chunking settings
CHUNK_SIZE = 800     # ~800 chars per chunk
CHUNK_OVERLAP = 100
MIN_CHUNK_SIZE = 80


class QueryRequest(BaseModel):
    site_id: str
    query: str
    n_results: int = 8


def _validate_site_id(site_id: str) -> str:
    if not _SITE_ID_RE.fullmatch(site_id or ""):
        raise HTTPException(status_code=400, detail="Invalid site_id (must be 24-hex ObjectId)")
    return site_id


def _get_collection(site_id: str):
    """取得(或建立)該工地的 collection。site_id 已驗證過格式。"""
    col = _collections.get(site_id)
    if col is None:
        col = chroma_client.get_or_create_collection(
            name=f"site_{site_id}",
            metadata={"hnsw:space": "cosine", "embedding_model": get_current_model_name()},
        )
        _collections[site_id] = col
    return col


def _check_collections_model(model_name: str) -> None:
    """startup 時逐 collection 檢查 embedding 模型是否一致,不符者刪除。

    有 embeddings.py 的持久化檔,正常不會發生;發生時代表有人改了 env,
    刪掉的 collection 對應的文件需在 CSIS 介面重新上傳(ai_documents 狀態會 stale)。
    """
    try:
        for col in chroma_client.list_collections():
            name = getattr(col, "name", str(col))
            if not name.startswith("site_"):
                continue
            existing = chroma_client.get_collection(name)
            old_model = (existing.metadata or {}).get("embedding_model", "unknown")
            if old_model != model_name:
                logger.warning(
                    f"[RAG] Embedding model changed for {name}: '{old_model}' -> '{model_name}'. "
                    f"Deleting collection ({existing.count()} chunks) — re-upload required!"
                )
                chroma_client.delete_collection(name)
    except Exception as e:
        logger.error(f"[RAG] Collection model check failed: {e}")


@app.on_event("startup")
async def startup():
    global chroma_client, _reranker

    os.makedirs(CHROMA_DIR, exist_ok=True)

    # Pre-load embedding model
    get_embedding_model()

    # Pre-load cross-encoder reranker(~2.2GB,startup 一次載完)
    if ENABLE_RERANKER:
        try:
            from sentence_transformers import CrossEncoder
            logger.info(f"[RAG] Loading reranker: {RERANKER_MODEL}...")
            _reranker = CrossEncoder(RERANKER_MODEL, max_length=RERANK_MAX_LENGTH)
            logger.info("[RAG] Reranker loaded")
        except Exception as e:
            logger.error(f"[RAG] Reranker load failed (fallback to RRF-only ranking): {e}")
            _reranker = None

    logger.info(f"[RAG] Initializing ChromaDB at {CHROMA_DIR}...")
    chroma_client = chromadb.PersistentClient(path=CHROMA_DIR)
    _check_collections_model(get_current_model_name())
    logger.info(f"[RAG] Ready. {len(chroma_client.list_collections())} site collections.")


@app.get("/health")
async def health():
    n_sites = 0
    n_chunks = 0
    if chroma_client is not None:
        try:
            for col in chroma_client.list_collections():
                name = getattr(col, "name", str(col))
                if name.startswith("site_"):
                    n_sites += 1
                    n_chunks += chroma_client.get_collection(name).count()
        except Exception:
            pass
    return {
        "status": "ok",
        "service": "csis-rag",
        "embedding_model": get_current_model_name(),
        "sites": n_sites,
        "chunks": n_chunks,
        "retrieval": "hybrid (BM25 + dense, RRF)" if _BM25_AVAILABLE else "dense only",
        "reranker": RERANKER_MODEL if _reranker is not None else None,
    }


@app.post("/upload")
async def upload_document(file: UploadFile = File(...), site_id: str = Form(...)):
    """解析並索引一份文件(PDF/DOCX/XLSX/PPTX/TXT/MD)到指定工地的 pool。

    原始檔由 Bun 端存 GridFS,這裡只做解析 → chunking → embedding → 入庫。
    PDF/PPTX 的 chunk 帶 page metadata(citation 跳頁用)。
    """
    _validate_site_id(site_id)
    col = _get_collection(site_id)

    filename = file.filename or "unknown"
    content_bytes = await file.read()
    ext = os.path.splitext(filename)[1].lower()

    # 依檔案類型 dispatch:能拿到頁碼就拿,拿不到就 fallback 到無頁碼路徑。
    # 回傳 (list[(chunk_text, page_or_None)], page_count_or_None)
    if ext == ".pdf" and pdfplumber is not None:
        pages_chunks, page_count = _chunk_pdf_with_pages(content_bytes)
    elif ext == ".pptx" and Presentation is not None:
        pages_chunks, page_count = _chunk_pptx_with_slides(content_bytes)
    elif ext in MARKITDOWN_EXTENSIONS:
        text = _convert_with_markitdown(content_bytes, ext)
        pages_chunks = [(c, None) for c in _chunk_text(text, CHUNK_SIZE, CHUNK_OVERLAP)]
        page_count = None
    else:
        # .txt, .md, 其他純文字
        text = content_bytes.decode("utf-8", errors="replace")
        pages_chunks = [(c, None) for c in _chunk_text(text, CHUNK_SIZE, CHUNK_OVERLAP)]
        page_count = None

    # Filter empty chunks(page parser 可能對掃描頁回空字串)+ 圖面噪音過濾
    pages_chunks = [(c, p) for c, p in pages_chunks if c and c.strip()]
    pages_chunks = [(c, p) for c, p in pages_chunks if not _is_low_info(c)]

    # 文字抽取後仍無內容 → 掃描/圖片型 PDF,用 OCR 補救
    if not pages_chunks and ext == ".pdf":
        logger.info(f"[RAG] No extractable text from {filename}, attempting OCR fallback...")
        ocr_chunks, ocr_pages = _ocr_pdf_with_pages(content_bytes)
        pages_chunks = [(c, p) for c, p in ocr_chunks if c and c.strip() and not _is_low_info(c)]
        page_count = page_count or ocr_pages

    if not pages_chunks:
        return {"error": "Empty document (no extractable text)", "filename": filename}

    chunks = [c for c, _ in pages_chunks]
    # embed 時前綴文件名,讓向量帶「文件身份」→ 解決跨文件相似章節的張冠李戴。
    # documents 仍存原文(citation 預覽用,不污染)。
    doc_label = os.path.splitext(filename)[0]
    embed_inputs = [f"《{doc_label}》\n{c}" for c in chunks]
    embeddings = embed_texts(embed_inputs)

    doc_id = str(uuid.uuid4())[:8]
    ids = [f"{doc_id}_chunk_{i}" for i in range(len(chunks))]
    metadatas: list[dict] = []
    for i, (_, page) in enumerate(pages_chunks):
        meta: dict = {
            "filename": filename,
            "doc_id": doc_id,
            "chunk_index": i,
            "file_ext": ext,
        }
        # ChromaDB 不接受 None,所以只在有頁碼時才寫入 page key
        if page is not None:
            meta["page"] = page
        metadatas.append(meta)

    col.add(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    _build_bm25(site_id)  # 新增文件後重建該 site 的 BM25 索引

    return {
        "status": "ok",
        "filename": filename,
        "doc_id": doc_id,
        "chunks": len(chunks),
        "page_count": page_count,
        "has_page_info": any(p is not None for _, p in pages_chunks),
        "site_documents": col.count(),
    }


def _tokenize(text: str) -> list[str]:
    """jieba 檢索模式分詞,過濾空白 token。"""
    return [t for t in jieba.lcut_for_search(text or "") if t.strip()]


def _build_bm25(site_id: str) -> None:
    """從該 site collection 的全部 documents 重建 BM25 索引(原文,不含《檔名》前綴)。"""
    if not _BM25_AVAILABLE:
        _bm25_by_site.pop(site_id, None)
        return
    try:
        col = _get_collection(site_id)
        data = col.get(include=["documents"])
        ids = data.get("ids") or []
        docs = data.get("documents") or []
        if not ids:
            _bm25_by_site.pop(site_id, None)
            return
        corpus = [_tokenize(d or "") for d in docs]
        _bm25_by_site[site_id] = (BM25Okapi(corpus), ids)
        logger.info(f"[RAG] BM25 index built for site {site_id}: {len(ids)} chunks")
    except Exception as e:
        logger.error(f"[RAG] BM25 build failed for site {site_id} (fallback to dense-only): {e}")
        _bm25_by_site.pop(site_id, None)


def _get_bm25(site_id: str):
    """lazy 取得該 site 的 BM25 索引(重啟後首查時重建)。"""
    if site_id not in _bm25_by_site:
        _build_bm25(site_id)
    return _bm25_by_site.get(site_id, (None, []))


def _fallback_score(distance) -> float:
    # 無 reranker 時的相關度近似:cosine distance → 1-d(BM25-only 命中無距離,給中性 0.5)
    if distance is None:
        return 0.5
    return max(0.0, min(1.0, 1.0 - float(distance)))


def _hybrid_search(site_id: str, query: str, n_results: int) -> list[dict]:
    """Dense + BM25 雙路召回,RRF 融合,cross-encoder 重排。全程限定在該 site 的 pool。

    每筆結果帶 rerank_score(0~1,citation 相關度顯示用)。
    """
    col = _get_collection(site_id)
    total = col.count()
    n_results = min(n_results, total)
    if n_results <= 0:
        return []

    cand = min(HYBRID_CANDIDATES, total)
    info: dict[str, dict] = {}

    # 1) Dense 召回
    dense = col.query(query_embeddings=[embed_query(query)], n_results=cand)
    dense_ids = (dense.get("ids") or [[]])[0]
    dense_docs = (dense.get("documents") or [[]])[0]
    dense_metas = (dense.get("metadatas") or [[]])[0]
    dense_dists = (dense.get("distances") or [[]])[0]
    for i, cid in enumerate(dense_ids):
        info[cid] = {
            "text": dense_docs[i] if i < len(dense_docs) else "",
            "metadata": dense_metas[i] if i < len(dense_metas) else {},
            "distance": dense_dists[i] if i < len(dense_dists) else None,
        }

    # 2) BM25 召回(per-site 索引)
    bm25_ids: list[str] = []
    bm25, bm25_all_ids = _get_bm25(site_id)
    if bm25 is not None and bm25_all_ids:
        q_tokens = _tokenize(query)
        if q_tokens:
            scores = bm25.get_scores(q_tokens)
            order = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)
            for i in order[:cand]:
                if scores[i] <= 0:
                    break
                bm25_ids.append(bm25_all_ids[i])

    # 3) RRF 融合 — 取前 RERANK_POOL 筆當重排候選(無 reranker 時直接取前 n_results)
    rrf: dict[str, float] = {}
    for rank, cid in enumerate(dense_ids):
        rrf[cid] = rrf.get(cid, 0.0) + 1.0 / (RRF_K + rank + 1)
    for rank, cid in enumerate(bm25_ids):
        rrf[cid] = rrf.get(cid, 0.0) + 1.0 / (RRF_K + rank + 1)
    pool_size = max(RERANK_POOL, n_results) if _reranker is not None else n_results
    top_ids = sorted(rrf.keys(), key=lambda c: rrf[c], reverse=True)[:pool_size]

    # 4) 補齊只被 BM25 命中(不在 dense 結果)的 chunk 的 text/metadata
    missing = [cid for cid in top_ids if cid not in info]
    if missing:
        got = col.get(ids=missing, include=["documents", "metadatas"])
        g_ids = got.get("ids") or []
        g_docs = got.get("documents") or []
        g_metas = got.get("metadatas") or []
        for i, cid in enumerate(g_ids):
            info[cid] = {
                "text": g_docs[i] if i < len(g_docs) else "",
                "metadata": g_metas[i] if i < len(g_metas) else {},
                "distance": None,  # BM25-only 命中,無向量距離
            }

    candidates = [info[cid] for cid in top_ids if info.get(cid) and info[cid]["text"]]

    # 5) Cross-encoder 重排 + rerank_score(sigmoid 到 0~1)。
    #    失敗時保底退回 RRF 排序,分數用 1-distance 近似。
    if _reranker is not None and candidates:
        try:
            import math
            scores = _reranker.predict([(query, c["text"]) for c in candidates])
            for i, c in enumerate(candidates):
                c["rerank_score"] = 1.0 / (1.0 + math.exp(-float(scores[i])))
            candidates.sort(key=lambda c: c["rerank_score"], reverse=True)
        except Exception as e:
            logger.error(f"[RAG] Rerank failed (fallback to RRF order): {e}")
            for c in candidates:
                c["rerank_score"] = _fallback_score(c.get("distance"))
    else:
        for c in candidates:
            c["rerank_score"] = _fallback_score(c.get("distance"))

    return candidates[:n_results]


@app.post("/query")
async def query_knowledge(req: QueryRequest):
    """查詢指定工地的知識庫 — Hybrid(dense + BM25,RRF 融合 + 重排)。"""
    _validate_site_id(req.site_id)
    col = _get_collection(req.site_id)
    if col.count() == 0:
        return {"results": [], "query": req.query}
    return {"results": _hybrid_search(req.site_id, req.query, req.n_results), "query": req.query}


@app.get("/documents")
async def list_documents(site_id: str):
    """列出指定工地已索引的文件(除錯用;正式列表以 CSIS 的 ai_documents 為準)。"""
    _validate_site_id(site_id)
    col = _get_collection(site_id)
    if col.count() == 0:
        return {"documents": []}

    all_data = col.get(include=["metadatas"])
    docs_map: dict[str, dict] = {}
    for meta in all_data["metadatas"]:
        doc_id = meta.get("doc_id", "unknown")
        if doc_id not in docs_map:
            docs_map[doc_id] = {
                "doc_id": doc_id,
                "filename": meta.get("filename", "unknown"),
                "chunks": 0,
            }
        docs_map[doc_id]["chunks"] += 1

    return {"documents": list(docs_map.values())}


@app.delete("/documents/{site_id}/{doc_id}")
async def delete_document(site_id: str, doc_id: str):
    """刪除指定工地中某文件的全部 chunks。"""
    _validate_site_id(site_id)
    if not re.fullmatch(r"[0-9a-f]{8}", doc_id):
        raise HTTPException(status_code=400, detail="Invalid doc_id format")

    col = _get_collection(site_id)
    matches = col.get(where={"doc_id": doc_id}, include=["metadatas"])
    ids_to_delete = matches.get("ids") or []
    if ids_to_delete:
        col.delete(ids=ids_to_delete)
        _build_bm25(site_id)

    return {
        "status": "ok",
        "doc_id": doc_id,
        "deleted_chunks": len(ids_to_delete),
        "site_documents": col.count(),
    }


def _is_low_info(text: str) -> bool:
    """判斷 chunk 是否為圖面噪音(軸線編號、座標、純數字序列)。

    工程文件的圖面頁常被 pdfplumber 抽成「1 2 3 ... A A B B C C」這類無語意
    內容,embed 後向量沒有語意卻會佔據檢索名額。保守判斷:只要有足夠中文字或
    英文單詞就保留,否則看數字/孤立單字母佔比是否過高。
    """
    t = text.strip()
    # 短的圖說/附件/頁眉標題:無實質內容,卻因檔名前綴佔據檢索前排
    if len(t) <= 45 and re.search(
        r"(設計圖|平面圖|立面圖|剖面圖|配置圖|示意圖|大樣圖|附件\s*[一二三四五六七八九十\d])", t
    ):
        return True
    cjk = len(re.findall("[一-鿿]", t))                     # 中文字數
    words = len(re.findall(r"[A-Za-z]{3,}", t))             # 有意義英文單詞(>=3 字母)
    if cjk >= 8 or words >= 5:
        return False
    non_space = re.sub(r"\s", "", t)
    if not non_space:
        return True
    noise = len(re.findall(r"[0-9]", non_space)) + len(
        re.findall(r"(?<![A-Za-z])[A-Za-z](?![A-Za-z])", non_space)
    )                                                       # 數字 + 孤立單字母
    return noise / len(non_space) > 0.6


def _convert_with_markitdown(file_bytes: bytes, ext: str) -> str:
    """Convert a binary document (PDF/DOCX/XLSX/PPTX/etc.) to Markdown text via markitdown."""
    try:
        result = markitdown.convert_stream(io.BytesIO(file_bytes), file_extension=ext)
        return result.text_content or ""
    except Exception as e:
        logger.error(f"[RAG] markitdown conversion failed for {ext}: {e}")
        return f"[Error converting document: {e}]"


_ocr_engine = None  # RapidOCR 單例(首次用到才初始化,載入 onnx 模型較慢)


def _ocr_pdf_with_pages(pdf_bytes: bytes) -> tuple[list[tuple[str, int | None]], int | None]:
    """對掃描/圖片型 PDF 做 OCR — pdfplumber + markitdown 都抽不到文字時的最後手段。

    用 PyMuPDF 把每頁 render 成點陣圖,再用 RapidOCR(onnxruntime,純 CPU)辨識。
    CPU OCR 很慢(每頁約數秒),屬離線上傳可接受範圍。
    依賴缺失或初始化失敗時回空(caller 會回報 Empty document)。
    """
    global _ocr_engine
    try:
        import fitz  # PyMuPDF
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
    except Exception as e:
        logger.error(f"[RAG] OCR dependencies unavailable: {e}")
        return [], None

    if _ocr_engine is None:
        logger.info("[RAG] Initializing RapidOCR engine (first use)...")
        _ocr_engine = RapidOCR()

    results: list[tuple[str, int | None]] = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        page_count = len(doc)
        for page_no in range(page_count):
            pix = doc[page_no].get_pixmap(dpi=200, colorspace=fitz.csRGB)
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            ocr_out, _ = _ocr_engine(img)
            if not ocr_out:
                continue
            page_text = "\n".join(line[1] for line in ocr_out).strip()
            if not page_text:
                continue
            for c in _chunk_text(page_text, CHUNK_SIZE, CHUNK_OVERLAP):
                results.append((c, page_no + 1))
        logger.info(f"[RAG] OCR done: {page_count} pages -> {len(results)} chunks")
    finally:
        doc.close()
    return results, page_count


def _extract_tables_md(page) -> list[str]:
    """把 pdfplumber 在該頁偵測到的表格各轉成一段 Markdown 表格。

    工程文件的關鍵數值幾乎都在表格;page.extract_text() 會把表格攤平成一串數字,
    LLM 無法對應欄位。轉成 Markdown 表格後每個數值都帶著欄標題。每個表格獨立成
    chunk(不經 _chunk_text 腰斬),維持欄列完整。
    """
    out: list[str] = []
    try:
        tables = page.extract_tables() or []
    except Exception:
        return out
    for tbl in tables:
        rows = [r for r in tbl if r and any((c or "").strip() for c in r)]
        if len(rows) < 2:
            continue                                  # 至少表頭 + 一列資料才有意義
        ncol = max(len(r) for r in rows)
        md: list[str] = []
        for i, r in enumerate(rows):
            cells = [((c or "").replace("\n", " ").strip()) for c in r]
            cells += [""] * (ncol - len(cells))       # 補齊不足的欄
            md.append("| " + " | ".join(cells) + " |")
            if i == 0:
                md.append("| " + " | ".join(["---"] * ncol) + " |")
        out.append("\n".join(md))
    return out


def _chunk_pdf_with_pages(pdf_bytes: bytes) -> tuple[list[tuple[str, int | None]], int | None]:
    """以頁為單位拆 PDF,每個 chunk 帶上頁碼(從 1 起算),並回傳總頁數。

    若 pdfplumber 因 PDF 結構毀損/加密等可預期原因失敗,退回到無頁碼的
    markitdown 路徑;其他 unexpected exception 直接 propagate,不可吞。
    """
    EXPECTED_PARSE_ERRORS = (ValueError, OSError, KeyError, IndexError)
    try:
        results: list[tuple[str, int | None]] = []
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            page_count = len(pdf.pages)
            for page_no, page in enumerate(pdf.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    for c in _chunk_text(page_text, CHUNK_SIZE, CHUNK_OVERLAP):
                        results.append((c, page_no))
                # 表格各自成獨立 chunk,保留欄列結構
                for tbl_md in _extract_tables_md(page):
                    results.append((tbl_md, page_no))
        if results:
            return results, page_count
        # pdfplumber 開得起來但拿不到任何文字(掃描型 PDF)→ caller 會走 OCR
        logger.warning("[RAG] pdfplumber extracted no text (likely scanned/image-only PDF) — falling back to markitdown")
    except EXPECTED_PARSE_ERRORS as e:
        logger.error(f"[RAG] pdfplumber parse failed ({type(e).__name__}): {e} — falling back to markitdown")
        page_count = None

    text = _convert_with_markitdown(pdf_bytes, ".pdf")
    return [(c, None) for c in _chunk_text(text, CHUNK_SIZE, CHUNK_OVERLAP)], page_count


def _chunk_pptx_with_slides(pptx_bytes: bytes) -> tuple[list[tuple[str, int | None]], int | None]:
    """以 slide 為單位拆 PPTX,page 對應 slide number(從 1 起算),回傳 slide 總數。

    抽出每個 slide 的文字框 + speaker notes(若有)。
    """
    EXPECTED_PARSE_ERRORS = (ValueError, OSError, KeyError, IndexError, AttributeError)
    slide_count: int | None = None
    try:
        results: list[tuple[str, int | None]] = []
        prs = Presentation(io.BytesIO(pptx_bytes))
        slides = list(prs.slides)
        slide_count = len(slides)
        for slide_no, slide in enumerate(slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
            try:
                notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
                if notes:
                    parts.append(f"[Speaker notes]\n{notes}")
            except AttributeError as e:
                logger.debug(f"[RAG] PPTX slide {slide_no} notes parse skipped: {e}")

            slide_text = "\n\n".join(p for p in parts if p.strip())
            if not slide_text.strip():
                continue
            for c in _chunk_text(slide_text, CHUNK_SIZE, CHUNK_OVERLAP):
                results.append((c, slide_no))
        if results:
            return results, slide_count
        logger.warning("[RAG] python-pptx extracted no text — falling back to markitdown")
    except EXPECTED_PARSE_ERRORS as e:
        logger.error(f"[RAG] python-pptx parse failed ({type(e).__name__}): {e} — falling back to markitdown")

    text = _convert_with_markitdown(pptx_bytes, ".pptx")
    return [(c, None) for c in _chunk_text(text, CHUNK_SIZE, CHUNK_OVERLAP)], slide_count


def _chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Markdown-aware text chunking.

    1. 依 Markdown 標題切 section,chunk 前綴 heading 麵包屑
    2. section 內按段落/表格/code fence 區塊累積到 chunk_size
    3. 過長區塊按句子邊界切,最後手段按字元切
    4. 碎片(< MIN_CHUNK_SIZE)併入鄰居
    """
    sections = _split_by_headings(text)
    chunks: list[str] = []

    for heading_ctx, body in sections:
        prefix = heading_ctx + "\n" if heading_ctx else ""
        prefix_len = len(prefix)
        effective_size = chunk_size - prefix_len

        if not body.strip():
            continue

        if len(body) <= effective_size:
            chunks.append(prefix + body.strip())
            continue

        sub_chunks = _split_body(body, effective_size, overlap)
        for sc in sub_chunks:
            chunks.append(prefix + sc)

    merged = _merge_small_chunks(chunks, MIN_CHUNK_SIZE, chunk_size)
    return merged if merged else [text[:chunk_size]]


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)


def _split_by_headings(text: str) -> list[tuple[str, str]]:
    """Split Markdown text into (heading_context, body) pairs.

    heading_context 是父層標題麵包屑,如 "# 手冊 > ## 退貨政策"。
    """
    matches = list(_HEADING_RE.finditer(text))

    if not matches:
        return [("", text)]

    sections: list[tuple[str, str]] = []
    heading_stack: list[tuple[int, str]] = []

    pre_text = text[: matches[0].start()].strip()
    if pre_text:
        sections.append(("", pre_text))

    for i, m in enumerate(matches):
        level = len(m.group(1))
        heading_text = m.group(0).strip()

        while heading_stack and heading_stack[-1][0] >= level:
            heading_stack.pop()
        heading_stack.append((level, heading_text))

        heading_ctx = " > ".join(h[1] for h in heading_stack)

        body_start = m.end()
        body_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = text[body_start:body_end].strip()

        sections.append((heading_ctx, body))

    return sections


def _split_body(body: str, max_size: int, overlap: int) -> list[str]:
    """Split a section body into chunks, respecting block boundaries."""
    blocks = _split_into_blocks(body)

    chunks: list[str] = []
    current = ""

    for block in blocks:
        if len(current) + len(block) + 2 <= max_size:
            current += ("\n\n" if current else "") + block
        else:
            if current:
                chunks.append(current.strip())
            if len(block) > max_size:
                for sub in _split_long_block(block, max_size, overlap):
                    chunks.append(sub)
                current = ""
            else:
                if chunks and overlap > 0:
                    prev_tail = _safe_overlap(chunks[-1], overlap)
                    current = prev_tail + "\n\n" + block
                else:
                    current = block

    if current.strip():
        chunks.append(current.strip())

    return chunks


def _split_into_blocks(text: str) -> list[str]:
    """Split text into semantic blocks: paragraphs, tables, code fences."""
    blocks: list[str] = []
    current_block_lines: list[str] = []
    in_code_fence = False
    in_table = False

    for line in text.split("\n"):
        if line.strip().startswith("```"):
            if in_code_fence:
                current_block_lines.append(line)
                blocks.append("\n".join(current_block_lines))
                current_block_lines = []
                in_code_fence = False
                continue
            else:
                if current_block_lines:
                    blocks.append("\n".join(current_block_lines))
                    current_block_lines = []
                current_block_lines.append(line)
                in_code_fence = True
                continue

        if in_code_fence:
            current_block_lines.append(line)
            continue

        is_table_line = line.strip().startswith("|") or re.match(r"^\s*\|?[\s\-:]+\|", line)
        if is_table_line:
            if not in_table and current_block_lines:
                blocks.append("\n".join(current_block_lines))
                current_block_lines = []
            in_table = True
            current_block_lines.append(line)
            continue
        elif in_table:
            blocks.append("\n".join(current_block_lines))
            current_block_lines = []
            in_table = False

        if not line.strip():
            if current_block_lines:
                blocks.append("\n".join(current_block_lines))
                current_block_lines = []
        else:
            current_block_lines.append(line)

    if current_block_lines:
        blocks.append("\n".join(current_block_lines))

    return [b.strip() for b in blocks if b.strip()]


def _split_long_block(block: str, max_size: int, overlap: int) -> list[str]:
    """Split an oversized block by sentences, then by characters as fallback."""
    sentences = re.split(r"(?<=[。！？；\.\!\?\;])\s*", block)
    if len(sentences) <= 1:
        return _char_split(block, max_size, overlap)

    chunks: list[str] = []
    current = ""
    for sent in sentences:
        if not sent.strip():
            continue
        if len(current) + len(sent) + 1 <= max_size:
            current += sent
        else:
            if current:
                chunks.append(current.strip())
            if len(sent) > max_size:
                chunks.extend(_char_split(sent, max_size, overlap))
                current = ""
            else:
                current = sent
    if current.strip():
        chunks.append(current.strip())
    return chunks


def _char_split(text: str, max_size: int, overlap: int) -> list[str]:
    """Last-resort character-level splitting."""
    chunks = []
    step = max_size - overlap
    for i in range(0, len(text), step):
        chunk = text[i : i + max_size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def _safe_overlap(text: str, overlap: int) -> str:
    """Extract overlap text, trying to break at a sentence/word boundary."""
    if len(text) <= overlap:
        return text
    tail = text[-overlap:]
    for sep in ["。", "！", "？", "；", ". ", "! ", "? "]:
        pos = tail.find(sep)
        if pos >= 0:
            return tail[pos + len(sep):]
    return tail


def _merge_small_chunks(chunks: list[str], min_size: int, max_size: int) -> list[str]:
    """Merge chunks smaller than min_size with their neighbor."""
    if not chunks:
        return chunks

    merged: list[str] = []
    for chunk in chunks:
        if merged and len(merged[-1]) < min_size and len(merged[-1]) + len(chunk) + 2 <= max_size:
            merged[-1] += "\n\n" + chunk
        elif merged and len(chunk) < min_size and len(merged[-1]) + len(chunk) + 2 <= max_size:
            merged[-1] += "\n\n" + chunk
        else:
            merged.append(chunk)
    return merged


if __name__ == "__main__":
    port = int(os.environ.get("RAG_PORT", "8010"))
    host = os.environ.get("RAG_HOST", "127.0.0.1")  # 只供本機 Bun 端呼叫,不對外
    logger.info(f"[RAG] Starting CSIS RAG worker on {host}:{port}...")
    uvicorn.run(app, host=host, port=port)
